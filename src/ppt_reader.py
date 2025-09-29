"""
PPT文件读取器
"""
import os
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional
from pptx import Presentation
from pdf2image import convert_from_path

logger = logging.getLogger(__name__)


class PPTReader:
    """PPT文件读取器"""
    
    def __init__(self):
        self.supported_formats = ['.pptx', '.ppt']
    
    def read_ppt(self, ppt_file: str) -> Optional[Dict[str, Any]]:
        """
        读取PPT文件和对应的PDF文件
        
        Args:
            ppt_file: PPT文件路径
            
        Returns:
            包含PPT和PDF内容的字典
        """
        try:
            ppt_path = Path(ppt_file)
            if not ppt_path.exists():
                logger.error(f"PPT文件不存在: {ppt_file}")
                return None
            
            # 查找对应的PDF文件
            pdf_file = self._find_pdf_file(ppt_path)
            if not pdf_file:
                logger.error(f"未找到对应的PDF文件: {ppt_file}")
                return None
            
            # 读取PPT内容
            ppt_content = self._read_ppt_content(ppt_path)
            
            # 读取PDF页面
            pdf_pages = self._read_pdf_pages(pdf_file)
            
            # 获取PDF基本信息
            pdf_info = self._get_pdf_info(pdf_file)
            
            return {
                'ppt_file': str(ppt_path),
                'pdf_file': pdf_file,
                'ppt_content': ppt_content,
                'pdf_pages': pdf_pages,
                'pdf_info': pdf_info,
                'total_pages': len(ppt_content),
                'file_info': self._get_file_info(ppt_path, Path(pdf_file))
            }
            
        except Exception as e:
            logger.error(f"读取PPT文件失败: {e}")
            return None
    
    def _find_pdf_file(self, ppt_path: Path) -> Optional[str]:
        """查找同名的PDF文件"""
        pdf_path = ppt_path.with_suffix('.pdf')
        if pdf_path.exists():
            return str(pdf_path)
        return None
    
    def _read_ppt_content(self, ppt_path: Path) -> List[Dict[str, Any]]:
        """读取PPT内容，提取每页的备注"""
        content = []
        
        try:
            prs = Presentation(ppt_path)
            
            for i, slide in enumerate(prs.slides):
                page_info = {
                    'page_number': i + 1,
                    'notes_text': '',
                    'is_full_mode': False
                }
                
                # 提取备注
                if slide.has_notes_slide:
                    notes_slide = slide.notes_slide
                    if notes_slide.notes_text_frame:
                        notes_text = notes_slide.notes_text_frame.text.strip()
                        page_info['notes_text'] = notes_text
                        
                        # 检查是否为full模式
                        if notes_text.startswith('[full]'):
                            page_info['is_full_mode'] = True
                            page_info['notes_text'] = notes_text[6:].strip()  # 移除[full]前缀
                
                content.append(page_info)
                
        except Exception as e:
            logger.error(f"读取PPT内容失败: {e}")
        
        return content
    
    def _read_pdf_pages(self, pdf_file: str) -> List[str]:
        """读取PDF文件，转换为图片路径列表"""
        pages = []
        
        try:
            # 转换PDF为图片
            images = convert_from_path(pdf_file)
            
            for i, image in enumerate(images):
                # 保存图片到临时目录
                temp_dir = Path("temp")
                temp_dir.mkdir(exist_ok=True)
                
                image_path = temp_dir / f"page_{i+1}.png"
                image.save(image_path, 'PNG')
                
                pages.append(str(image_path))
                
        except Exception as e:
            logger.error(f"读取PDF页面失败: {e}")
        
        return pages
    
    def _get_pdf_info(self, pdf_file: str) -> Dict[str, Any]:
        """获取PDF文件基本信息"""
        pdf_info = {
            'total_pages': 0,
            'file_size': 0,
            'file_exists': False,
            'is_readable': False,
            'conversion_successful': False
        }
        
        try:
            pdf_path = Path(pdf_file)
            if not pdf_path.exists():
                logger.warning(f"PDF文件不存在: {pdf_file}")
                return pdf_info
            
            pdf_info['file_exists'] = True
            pdf_info['file_size'] = pdf_path.stat().st_size
            
            # 尝试读取PDF页面数量
            images = convert_from_path(pdf_file)
            pdf_info['total_pages'] = len(images)
            pdf_info['is_readable'] = True
            pdf_info['conversion_successful'] = True
            
            logger.info(f"PDF基本信息 - 总页数: {pdf_info['total_pages']}, "
                       f"文件大小: {pdf_info['file_size']} 字节")
            
        except Exception as e:
            logger.error(f"获取PDF信息失败: {e}")
            pdf_info['is_readable'] = False
        
        return pdf_info
    
    def _get_file_info(self, ppt_path: Path, pdf_path: Path) -> Dict[str, Any]:
        """获取文件信息"""
        file_info = {
            'ppt_file': {
                'name': ppt_path.name,
                'size': ppt_path.stat().st_size if ppt_path.exists() else 0,
                'created': ppt_path.stat().st_ctime if ppt_path.exists() else 0,
                'modified': ppt_path.stat().st_mtime if ppt_path.exists() else 0
            },
            'pdf_file': {
                'name': pdf_path.name,
                'size': pdf_path.stat().st_size if pdf_path.exists() else 0,
                'created': pdf_path.stat().st_ctime if pdf_path.exists() else 0,
                'modified': pdf_path.stat().st_mtime if pdf_path.exists() else 0
            }
        }
        return file_info